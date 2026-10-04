import pytest
from docx import Document
from pptx import Presentation

from app.parsers import parse_file


def test_txt(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("hello world")
    assert parse_file(str(p)) == "hello world"


def test_docx(tmp_path):
    p = tmp_path / "a.docx"
    d = Document()
    d.add_paragraph("Cells have a nucleus.")
    d.save(str(p))
    assert "nucleus" in parse_file(str(p))


def test_pptx(tmp_path):
    p = tmp_path / "a.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Mitochondria"
    prs.save(str(p))
    assert "Mitochondria" in parse_file(str(p))


def test_demo_pdf_is_parseable():
    assert len(parse_file("data/demo/demo_doc.pdf")) > 100


def test_unsupported_extension(tmp_path):
    p = tmp_path / "a.exe"
    p.write_bytes(b"x")
    with pytest.raises(ValueError):
        parse_file(str(p))


def test_empty_file_rejected(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("   ")
    with pytest.raises(ValueError):
        parse_file(str(p))
