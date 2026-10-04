from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


def parse_pdf(path: str) -> str:
    reader = PdfReader(path)
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def parse_docx(path: str) -> str:
    doc = Document(path)
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def parse_pptx(path: str) -> str:
    out = []
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
    return "\n\n".join(out)


def parse_text(path: str) -> str:
    return Path(path).read_text(encoding="utf-8", errors="ignore")


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".pptx": parse_pptx,
    ".txt": parse_text,
    ".md": parse_text,
}


def parse_file(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext not in PARSERS:
        raise ValueError(f"Unsupported file type: {ext}")
    text = PARSERS[ext](path)
    if not text.strip():
        raise ValueError("No extractable text (scanned PDF?)")
    return text
